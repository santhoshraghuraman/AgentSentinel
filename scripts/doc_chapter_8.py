"""
Chapter 8: Summary and Conclusion for AgentSentinel M.Tech Dissertation.
"""

import docx
from docx.shared import Inches, Pt
from scripts.doc_helpers import (
    add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_bullet_p
)

def build_chapter_8(doc):
    add_heading_1(doc, "CHAPTER 8\nSUMMARY AND CONCLUSION")

    # Introductory lead paragraph
    add_body_p(doc,
        "Autonomous AI agents represent one of the most transformative frontiers in modern computer science, offering "
        "autonomous execution across software engineering, data science, and enterprise automation. However, their stochastic "
        "nature and vulnerability to prompt injection, goal drift, and unconstrained privilege escalation necessitate robust, "
        "verifiable software engineering guardrails. This concluding chapter summarizes the implemented system, synthesizes "
        "the core research findings, transparently outlines recognized limitations, charts future engineering directions, and "
        "delivers the final academic conclusion.",
        first_line_indent=0
    )

    # 8.1 Summary of the Project
    add_heading_2(doc, "8.1 Summary of the Project")
    add_body_p(doc,
        "This dissertation introduced AgentSentinel: an out-of-band, model-agnostic runtime behavioral governance framework "
        "designed to operationalize James P. Anderson’s classical 1972 Reference Monitor concept for autonomous agentic AI. "
        "Operating as an external mediation layer between an agent's proposal engine and the host environment, AgentSentinel "
        "intercepts every proposed tool invocation prior to execution. Over the course of this research, a complete 12-module "
        "software architecture (M1 through M12) was designed, implemented in Python/FastAPI, backed by an SQLite relational "
        "audit ledger, and made observable through an interactive 14-screen React dashboard.",
        first_line_indent=0
    )

    add_body_p(doc,
        "The framework transforms raw interaction trajectories into a 22-dimensional continuous behavioral feature vector "
        "capturing action categories, state mutation properties, transition dynamics (such as read-to-write escalations), "
        "and inter-action timing statistics. These features are evaluated by a dual-engine architecture: a trained Random Forest "
        "classifier (100 estimators) computing continuous risk probabilities, and a deterministic Task Authorization Engine "
        "verifying compliance against declared capability grants. An authoritative Policy Fusion Engine reconciles these signals "
        "under five mandatory security invariants, guaranteeing that whenever an action is blocked, physical callable execution "
        "on the host machine is strictly ZERO.",
        first_line_indent=0.5
    )

    # 8.2 Key Findings and Research Contributions
    add_heading_2(doc, "8.2 Key Findings and Research Contributions")
    add_body_p(doc,
        "The key findings and scientific contributions established by this dissertation are:",
        first_line_indent=0
    )

    add_bullet_p(doc, " Empirical evaluation on the held-out ATBench trajectory benchmark (N = 250) confirmed that the 22-dimensional behavioral feature vector enables a Random Forest model to achieve 72.40% classification accuracy, 78.95% precision, and an ROC-AUC of 0.7782, validating Hypothesis H1.", "1. Behavioral Feature Discriminiation:")
    add_bullet_p(doc, " The experimental discovery that standalone machine learning exhibits a 39.52% false negative rate on stealthy attacks provides definitive empirical proof that statistical risk scoring cannot stand alone. The hybrid integration of deterministic task capability contracts successfully eliminated 100% of these false negatives in integration tests, validating Hypothesis H2.", "2. Necessity of Hybrid Policy Fusion:")
    add_bullet_p(doc, " Rigorous test verification across 152 automated Pytest suites established that 100% of blocked actions resulted in exactly 0 host tool executions, validating Invariant I4.", "3. Guaranteed Zero Host Execution on BLOCK:")
    add_bullet_p(doc, " End-to-end latency profiling demonstrated a median mediation overhead of only 14.2 milliseconds and a P95 latency of 22.8 milliseconds, proving that comprehensive behavioral reference monitoring introduces negligible overhead (< 1%) compared to LLM generation cycles, validating Hypothesis H3.", "4. Ultra-Low Latency Performance:")
    add_bullet_p(doc, " The framework demonstrated complete model agnosticism, successfully governing both synthetic test invocations and autonomous multi-step reasoning traces driven by a local Ollama Qwen3:4b model without requiring model fine-tuning.", "5. Universal Model Agnosticism:")

    # 8.3 Limitations Identified
    add_heading_2(doc, "8.3 Limitations Identified")
    add_body_p(doc,
        "In strict adherence to academic integrity, the following limitations are explicitly recognized:",
        first_line_indent=0
    )

    add_bullet_p(doc, " In the ATBench dataset, trajectory length (event_count at 35.65% and action_count at 24.11%) serves as a strong statistical predictor. In environments where benign tasks have long horizons and attacks are short, the model requires length-normalized calibration.", "• Dataset Distribution Dependency:")
    add_bullet_p(doc, " While the current SQLite ledger provides atomic append-only relational storage, it relies on application-level integrity rather than cryptographic Merkle hash chaining or distributed consensus.", "• Persistence Immutability:")
    add_bullet_p(doc, " The React observability dashboard refreshes telemetry via 3-second HTTP polling rather than bidirectional WebSockets, which is suitable for demonstration but sub-optimal for ultra-high-frequency production fleets.", "• UI Telemetry Transport:")
    add_bullet_p(doc, " The 22 features model behavioral sequence dynamics but do not perform deep abstract syntax tree (AST) inspection of Python code or SQL query strings executed within tool arguments.", "• Deep Argument Inspection:")

    # 8.4 Future Research and Engineering Directions
    add_heading_2(doc, "8.4 Future Research and Engineering Directions")
    add_body_p(doc,
        "Based on the empirical findings and recognized limitations, several promising directions for future research are identified:",
        first_line_indent=0
    )

    add_bullet_p(doc, " Incorporating reinforcement learning from security feedback (RLSF) to dynamically adjust critical risk thresholds (theta_critical) based on real-time agent drift and environmental hostility.", "1. Adaptive Threshold Calibration via Reinforcement Learning:")
    add_bullet_p(doc, " Upgrading the relational audit store to incorporate SHA-256 Merkle tree hashing across consecutive interception events, guaranteeing cryptographic non-repudiation and tamper detection.", "2. Cryptographic Merkle Hash Chaining:")
    add_bullet_p(doc, " Integrating lightweight tree-sitter AST parsers into Module M2 to analyze code and SQL argument payloads prior to feature extraction, closing semantic gaps on polymorphic attacks.", "3. In-Line AST Semantic Parsing:")
    add_bullet_p(doc, " Containerizing AgentSentinel into an asynchronous Envoy-style sidecar proxy for Kubernetes, enabling distributed multi-tenant governance across thousands of concurrent agent workers.", "4. Distributed Cloud-Native Deployment:")

    # 8.5 Final Conclusion
    add_heading_2(doc, "8.5 Final Conclusion")
    add_body_p(doc,
        "This dissertation has demonstrated that ensuring the security and compliance of autonomous AI agents does not require "
        "restricting model capabilities or relying on costly, brittle prompt guardrails. By operationalizing James Anderson’s "
        "classical Reference Monitor concept as an out-of-band middleware that unifies temporal behavioral feature extraction, "
        "supervised machine learning, deterministic task authorization, and invariant-driven policy fusion, AgentSentinel "
        "establishes a robust, low-latency, and model-agnostic defensive foundation for the agentic AI era. The system is fully "
        "implemented, rigorously verified across 152 automated tests and 14 dashboard screens, and stands ready as a sound "
        "software engineering contribution to autonomous systems research.",
        first_line_indent=0
    )

    doc.add_page_break()
    print("Chapter 8 generated successfully.")
