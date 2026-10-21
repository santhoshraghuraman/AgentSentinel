"""
Chapter 6: Dissertation Demonstration for AgentSentinel M.Tech Dissertation.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from scripts.doc_helpers import (
    add_heading_1, add_heading_2, add_heading_3,
    add_body_p, add_bullet_p, add_figure
)

def build_chapter_6(doc):
    add_heading_1(doc, "CHAPTER 6\nDISSERTATION DEMONSTRATION")

    # Introductory lead paragraph
    add_body_p(doc,
        "A critical milestone of this dissertation is the empirical, live demonstration of the implemented AgentSentinel "
        "framework. To prove that the reference monitor operates robustly in real-world software engineering environments, "
        "this chapter provides the complete startup runbook, walks through all 14 interactive observability screens, and details "
        "three representative execution scenarios: a benign file read, an adversarial privilege escalation attempt, and an autonomous "
        "multi-step task driven by a local LLM agent. Every screenshot included in this chapter is captured directly from the live, "
        "running AgentSentinel system.",
        first_line_indent=0
    )

    # 6.1 Demonstration Environment and Launch Runbook
    add_heading_2(doc, "6.1 Demonstration Environment and Launch Runbook")
    add_body_p(doc,
        "The demonstration environment executes on a standard local workstation running Windows PowerShell, Python 3.11/3.14, "
        "and Node.js v22. The system operates across three decoupled service processes launched in separate terminal sessions:",
        first_line_indent=0
    )

    add_bullet_p(doc, " Starts the local LLM inference engine hosting Qwen3:4b on localhost:11434 with a 4,096-token context window. Run: `ollama run qwen3:4b`.", "Terminal 1 (Agent Proposal Engine):")
    add_bullet_p(doc, " Activates the Python virtual environment (.venv) and launches the Uvicorn ASGI server hosting the AgentSentinel reference monitor on port 8000. Run: `python -m uvicorn src.agentsentinel.api.main:create_app --factory --port 8000`.", "Terminal 2 (FastAPI Governance Backend):")
    add_bullet_p(doc, " Launches the Vite development server in ui/ to serve the React 19 single-page application on port 5173. Run: `npm run dev`.", "Terminal 3 (React Observability Frontend):")

    add_body_p(doc,
        "Once initiated, the user navigates to `http://localhost:5173` to access the full 14-screen AgentSentinel governance console.",
        first_line_indent=0.5
    )

    # 6.2 System Overview (2D Metrics & 3D Topology)
    add_heading_2(doc, "6.2 System Overview (2D Metrics & 3D Topology)")
    add_body_p(doc,
        "The Overview Dashboard serves as the central security operations center. As depicted in Figure 6.1, the 2D overview "
        "displays high-level telemetry widgets: total intercepted events (285 historical interceptions), governance decisions "
        "(163 ALLOW, 47 ALERT, 75 BLOCK), active agent sessions, and the multi-axis Sentinel Radar assessing real-time fleet risk posture.",
        first_line_indent=0
    )

    add_figure(doc, "01_overview_2d.png", "6.1", "Overview Dashboard (2D View with Sentinel Radar & Real-Time Metrics)", width_inches=5.8)

    add_body_p(doc,
        "To provide intuitive architectural observability, the dashboard features a toggle to a dynamic 3D topological visualization. "
        "Figure 6.2 shows the 3D topology highlighting the Policy Fusion Engine (M7) and Enforcement Gate (M8) nodes. Users can "
        "interactively rotate the graph to inspect connection health between the agent harness, the feature pipeline, and SQLite persistence.",
        first_line_indent=0.5
    )

    add_figure(doc, "02_overview_3d_m7.png", "6.2", "System Overview 3D Architectural Topology (Policy & Gate Nodes)", width_inches=5.8)

    # 6.3 Live Telemetry Monitoring
    add_heading_2(doc, "6.3 Live Telemetry Monitoring")
    add_body_p(doc,
        "The Live Monitor screen (Figure 6.3) streams real-time telemetry from active agent sessions. Built using React Query with "
        "a 3,000-millisecond HTTP polling interval (`refetchInterval: 3000`), the monitor tracks active agents, session durations, "
        "sliding-window action frequencies, and instantaneous risk meters. Any anomalous spike in tool-calling velocity immediately "
        "triggers visual warning indicators.",
        first_line_indent=0
    )

    add_figure(doc, "03_live_monitor.png", "6.3", "Live Telemetry Stream with Real-Time Risk Gauges", width_inches=5.8)

    # 6.4 Runtime Playground Scenarios
    add_heading_2(doc, "6.4 Runtime Playground Scenarios (Safe, Adversarial, Autonomous)")
    add_body_p(doc,
        "The Runtime Playground (`/playground`) provides an interactive testbed for simulating, verifying, and demonstrating the "
        "reference monitor under diverse operational conditions. Figure 6.4a shows the initial playground configuration interface, "
        "where operators select the agent identity, active task contract, tool primitive, and JSON argument payload.",
        first_line_indent=0
    )

    add_figure(doc, "04a_playground_initial.png", "6.4a", "Runtime Interactive Playground Configuration View", width_inches=5.8)

    add_heading_3(doc, "Scenario A: Benign Safe Operation (ALLOW -> Executed)")
    add_body_p(doc,
        "In the first scenario, a developer agent proposes a harmless documentation retrieval action: `read_file` with path "
        "`docs/README.md`. As shown in Figure 6.4b, AgentSentinel intercepts the action, extracts the 22-dimensional feature vector, "
        "and computes a low ML risk score (R_ML = 0.08). The Task Authorization Engine confirms that `file_system_read` is permitted "
        "in the active contract. M7 issues an ALLOW verdict. The interceptor invokes the host callable, returns the file contents, "
        "and logs the execution state as `executed`.",
        first_line_indent=0
    )

    add_figure(doc, "04b_playground_allow_outcome.png", "6.4b", "Playground Safe Tool Execution Scenario (ALLOW -> Executed)", width_inches=5.8)

    add_heading_3(doc, "Scenario B: Adversarial Privilege Escalation (BLOCK -> Zero Execution)")
    add_body_p(doc,
        "In the second scenario, a compromised agent attempts a destructive credential overwrite: `write_file` targeting `/etc/passwd`. "
        "As shown in Figure 6.4c, the reference monitor intercepts the payload. The Task Authorization Engine identifies that write "
        "privileges are not granted for system files, returning DENY. Invariant I1 (Denial Dominance) forces an immediate BLOCK verdict. "
        "Crucially, the host file system is NEVER touched (physical callable execution is strictly 0). The interceptor returns a "
        "synthetic error: 'Action BLOCKED by Sentinel Policy: Unauthorized file write attempt', and commits the audit state as `blocked`.",
        first_line_indent=0
    )

    add_figure(doc, "04c_playground_block_outcome.png", "6.4c", "Playground Adversarial Attack Scenario (BLOCK -> Zero Execution)", width_inches=5.8)

    add_heading_3(doc, "Scenario C: Autonomous LLM Agent Execution")
    add_body_p(doc,
        "In the third scenario, an autonomous multi-step workflow is initiated using the local Qwen3:4b LLM via Ollama. Figure 6.4d "
        "illustrates the agent autonomously decomposing a complex task into discrete tool steps. Every step generated by Qwen3 is "
        "dynamically intercepted and mediated by AgentSentinel, proving seamless model-agnostic governance without human intervention.",
        first_line_indent=0
    )

    add_figure(doc, "04d_playground_autonomous_agent.png", "6.4d", "Autonomous Agent Execution Powered by Local Qwen3:4b LLM", width_inches=5.8)

    # 6.5 Decision Explorer & Feature Radar
    add_heading_2(doc, "6.5 Decision Explorer & Feature Radar")
    add_body_p(doc,
        "The Decision Explorer screen (Figure 6.5) enables forensic deep-dive analysis of individual interception decisions. "
        "Selecting an intercepted event renders a 22-feature spider radar chart comparing the action's feature values against benign "
        "population baselines, displays the exact Random Forest posterior probability, and shows the governing security invariant.",
        first_line_indent=0
    )

    add_figure(doc, "05_decision_explorer_event_selected.png", "6.5", "Decision Explorer Deep-Dive with 22-Feature Radar Chart", width_inches=5.8)

    # 6.6 Security Trajectories and Transition Graphs
    add_heading_2(doc, "6.6 Security Trajectories and Transition Graphs")
    add_body_p(doc,
        "Figure 6.6 presents the Security Trajectories screen (`/trajectories`), which reconstructs complete agent sessions as "
        "chronological timelines and state-transition graphs. Operators can inspect sequence escalation patterns, such as sequences "
        "where an agent reads multiple files before attempting an outbound network post.",
        first_line_indent=0
    )

    add_figure(doc, "06_security_trajectories_session.png", "6.6", "Security Trajectory Timeline and State Transition Graph", width_inches=5.8)

    # 6.7 Governance Policies and Task Authorization Management
    add_heading_2(doc, "6.7 Governance Policies and Task Authorization Management")
    add_body_p(doc,
        "Policy administration is decoupled from application code. Figure 6.7 shows the Policies screen (`/policies`), where "
        "administrators configure declarative rules, priority rankings, and action triggers. Figure 6.8 displays the Task Authorization "
        "registry (`/tasks`), detailing active capability grants across agents.",
        first_line_indent=0
    )

    add_figure(doc, "08_policies_governance.png", "6.7", "Active Governance Policy Rulebook and Priority Configuration", width_inches=5.8)
    add_figure(doc, "09_tasks_authorization.png", "6.8", "Task Contract Registry & Capability Allow/Deny Grids", width_inches=5.8)

    # 6.8 Forensic Audit Trail & Evidence Provenance
    add_heading_2(doc, "6.8 Forensic Audit Trail & Evidence Provenance")
    add_body_p(doc,
        "The Audit Logs screen (Figure 6.9) provides a non-repudiable forensic ledger of all historical interceptions. Operators "
        "can search by UUIDv4 `interception_id`, filter by decision (ALLOW/ALERT/BLOCK) or execution state (`attempted`, `blocked`, "
        "`executed`, `execution_failed`), and verify that argument parameters survived intact without truncation or credential leakage.",
        first_line_indent=0
    )

    add_figure(doc, "10_audit_logs.png", "6.9", "Forensic Audit Log Trail Searchable by Interception UUID", width_inches=5.8)

    # 6.9 Datasets, Evaluation Metrics & Compliance Reports
    add_heading_2(doc, "6.9 Datasets, Evaluation Metrics & Compliance Reports")
    add_body_p(doc,
        "Figures 6.10, 6.11, and 6.12 showcase the scientific research and compliance verification interfaces. Figure 6.10 displays "
        "the Dataset Explorer inspecting ATBench and SyFI partitions. Figure 6.11 renders model performance metrics, confusion matrices, "
        "and feature importance rankings. Figure 6.12 displays the automated Compliance Audit Report, which verifies policy adherence "
        "across enterprise governance standards.",
        first_line_indent=0
    )

    add_figure(doc, "11_datasets_research.png", "6.10", "Research Dataset Explorer (ATBench and SyFI Partitions)", width_inches=5.8)
    add_figure(doc, "12_evaluation_metrics.png", "6.11", "Model Performance Visualizer and Confusion Matrix Display", width_inches=5.8)
    add_figure(doc, "13_compliance_reports.png", "6.12", "System Compliance Audit & Security Posture Summary", width_inches=5.8)

    # 6.10 System Diagnostics and Service Health
    add_heading_2(doc, "6.10 System Diagnostics and Service Health")
    add_body_p(doc,
        "Finally, Figures 6.13 and 6.14 show the Settings and System Diagnostics screens. Figure 6.13 allows operators to adjust "
        "detection thresholds (e.g., tuning $\\theta_{critical}$) and toggle features. Figure 6.14 provides real-time health checks "
        "for the FastAPI gateway, SQLite persistence, and local Ollama service, ensuring transparent operational status.",
        first_line_indent=0
    )

    add_figure(doc, "14_settings_configuration.png", "6.13", "Runtime Settings, Detection Thresholds and Feature Toggles", width_inches=5.8)
    add_figure(doc, "15_system_diagnostics.png", "6.14", "System Diagnostics and Component Health Monitor", width_inches=5.8)

    doc.add_page_break()
    print("Chapter 6 generated successfully.")
